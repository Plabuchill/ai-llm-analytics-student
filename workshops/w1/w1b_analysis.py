"""W1-B: ยอดขายสุทธิ กำไรขั้นต้น AOV ยอดรายเดือน และสัดส่วนตามช่องทาง
ไตรมาส ต.ค.–ธ.ค. 2025 เทียบ ก.ค.–ก.ย. 2026 (นิยามตาม w4-analytics-starter/CLAUDE.md)
"""
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill

plt.rcParams["font.family"] = "Tahoma"

ROOT = Path(__file__).resolve().parents[2]  # root ของ repo (ใช้ data/ ของ repo)
OUT = Path(__file__).parent / "outputs"
OUT.mkdir(exist_ok=True)

OK = "จัดส่งสำเร็จ"
QUARTERS = {
    "Q ต.ค.–ธ.ค. 2025": ["2025-10", "2025-11", "2025-12"],
    "Q ก.ค.–ก.ย. 2026": ["2026-07", "2026-08", "2026-09"],
}

# ---------- 1) อ่านข้อมูลและตรวจสอบ ----------
df = pd.read_csv(ROOT / "data/orders.csv", encoding="utf-8-sig")
products = pd.read_csv(ROOT / "data/products.csv", encoding="utf-8-sig")

raw_shape = df.shape
print("df.shape:", raw_shape)
print("products.shape:", products.shape)
print("คอลัมน์ products ก่อนแก้:", list(products.columns))
df.columns = df.columns.str.strip()
products.columns = products.columns.str.strip()
# ราคา/ต้นทุนใน products.csv เป็นข้อความ เช่น " 1,290 " ต้องตัดช่องว่างและจุลภาคก่อนแปลงเป็นตัวเลข
for col in ["list_price", "unit_cost"]:
    products[col] = pd.to_numeric(products[col].astype(str).str.strip().str.replace(",", ""), errors="coerce")
    assert products[col].notna().all(), f"{col} มีค่าที่แปลงเป็นตัวเลขไม่ได้"
print("\nจำนวนแต่ละ status:")
print(df["status"].value_counts().to_string())
print("\norder_id ซ้ำ:", df["order_id"].duplicated().sum())

df["month"] = df["order_datetime"].str[:7]
df["net_sales"] = df["qty"] * df["unit_price"] * (1 - df["discount_pct"] / 100)

df = df.merge(products[["product_id", "unit_cost"]], on="product_id", how="left")
print("แถวที่หา unit_cost ไม่เจอ:", df["unit_cost"].isna().sum())
df["cogs"] = df["qty"] * df["unit_cost"]

all_status = df["net_sales"].sum()
ok = df[df["status"] == OK].copy()
correct = ok["net_sales"].sum()
print(f"\nยอดรวมนับทุกสถานะ (ผิด): {all_status:,.2f}")
print(f"ยอดขายสุทธิ (ถูก)      : {correct:,.2f}")
print(f"ส่วนต่าง               : {all_status - correct:,.2f}")


# ---------- 2) KPI รายไตรมาส ----------
def kpi(months):
    allq = df[df["month"].isin(months)]
    q = ok[ok["month"].isin(months)]
    net = q["net_sales"].sum()
    gp = net - q["cogs"].sum()
    return {
        "จำนวนคำสั่งซื้อทั้งหมด": len(allq),
        "คำสั่งซื้อสำเร็จ": q["order_id"].nunique(),
        "ยอดขายสุทธิ": net,
        "กำไรขั้นต้น": gp,
        "% กำไรขั้นต้น": gp / net * 100,
        "AOV": net / q["order_id"].nunique(),
        "อัตรายกเลิก %": (allq["status"] == "ยกเลิก").mean() * 100,
        "อัตราคืน %": (allq["status"] == "คืนสินค้า").mean() * 100,
    }


summary = pd.DataFrame({name: kpi(m) for name, m in QUARTERS.items()})
a, b = summary.columns
summary["เปลี่ยนแปลง"] = summary[b] - summary[a]
summary["% เปลี่ยนแปลง"] = summary["เปลี่ยนแปลง"] / summary[a] * 100
print("\nKPI รายไตรมาส:")
print(summary.round(2).to_string())

# ---------- 3) ยอดรายเดือน (ทั้งปีงบ) ----------
monthly = (
    ok.groupby("month")
    .agg(orders=("order_id", "nunique"), net=("net_sales", "sum"), cogs=("cogs", "sum"))
    .reset_index()
    # ตั้งชื่อไทยผ่าน rename: ชื่อ keyword ภาษาไทยถูก Python ทำ NFKC ทำให้สระอำเพี้ยน
    .rename(columns={"month": "เดือน", "orders": "คำสั่งซื้อสำเร็จ", "net": "ยอดขายสุทธิ", "cogs": "ต้นทุน"})
)
monthly["กำไรขั้นต้น"] = monthly["ยอดขายสุทธิ"] - monthly["ต้นทุน"]
monthly["AOV"] = monthly["ยอดขายสุทธิ"] / monthly["คำสั่งซื้อสำเร็จ"]
monthly = monthly.drop(columns="ต้นทุน")
print("\nยอดรายเดือน:")
print(monthly.round(2).to_string(index=False))

fig, ax = plt.subplots(figsize=(10, 4.5))
ax.plot(monthly["เดือน"], monthly["ยอดขายสุทธิ"] / 1e6, marker="o", color="#2563eb", label="ยอดขายสุทธิ")
ax.plot(monthly["เดือน"], monthly["กำไรขั้นต้น"] / 1e6, marker="o", color="#16a34a", label="กำไรขั้นต้น")
for months, color in zip(QUARTERS.values(), ["#fde68a", "#bfdbfe"]):
    idx = monthly.index[monthly["เดือน"].isin(months)]
    ax.axvspan(idx.min() - 0.4, idx.max() + 0.4, color=color, alpha=0.4)
ax.set_title("ยอดขายสุทธิและกำไรขั้นต้นรายเดือน (ล้านบาท) · แถบสี = ไตรมาสที่เทียบ")
ax.set_ylabel("ล้านบาท")
ax.set_ylim(bottom=0)
ax.grid(axis="y", alpha=0.3)
ax.legend()
plt.xticks(rotation=45)
plt.tight_layout()
chart_path = OUT / "w1b_monthly_sales.png"
fig.savefig(chart_path, dpi=120)

# ---------- 4) สัดส่วนยอดขายตามช่องทาง ----------
ok["quarter"] = None
for name, months in QUARTERS.items():
    ok.loc[ok["month"].isin(months), "quarter"] = name
channel = ok.dropna(subset=["quarter"]).pivot_table(
    index="channel", columns="quarter", values="net_sales", aggfunc="sum"
)[list(QUARTERS)]
share = channel / channel.sum() * 100
channel_table = pd.concat(
    {"ยอดขายสุทธิ": channel, "สัดส่วน %": share}, axis=1
).reset_index().rename(columns={"channel": "ช่องทาง"})
channel_table.columns = [" ".join(c).strip() if isinstance(c, tuple) else c for c in channel_table.columns]
print("\nสัดส่วนตามช่องทาง:")
print(channel_table.round(2).to_string(index=False))

# ---------- 5) ตรวจสอบ ----------
check = pd.DataFrame(
    {
        "รายการ": ["df.shape", "ยอดรวมนับทุกสถานะ (ผิด)", "ยอดขายสุทธิ (ถูก)", "ส่วนต่าง"]
        + [f"status: {s}" for s in df["status"].value_counts().index],
        "ค่า": [str(raw_shape), all_status, correct, all_status - correct]
        + list(df["status"].value_counts().values),
    }
)

# ---------- 6) Excel ----------
xlsx = OUT / "w1b_summary.xlsx"
with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
    summary.reset_index(names="ตัวชี้วัด").to_excel(w, sheet_name="สรุปไตรมาส", index=False)
    monthly.to_excel(w, sheet_name="รายเดือน", index=False)
    channel_table.to_excel(w, sheet_name="ช่องทาง", index=False)
    check.to_excel(w, sheet_name="ตรวจสอบ", index=False)

    header = PatternFill("solid", fgColor="DBEAFE")
    for ws in w.book.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(14, len(str(col[0].value)) + 4)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
    w.book["รายเดือน"].add_image(XLImage(chart_path), "H2")

print(f"\nบันทึกแล้ว: {xlsx}\nกราฟ: {chart_path}")
